from flask import Flask, jsonify, request, render_template_string, send_file
import io
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
import asyncio
import re
from playwright.async_api import async_playwright

app = Flask(__name__)
app.json.sort_keys = False

COOKIE_STRING = """
    datr=0f4TaniQolJ5C4YoG81CwSqZ;
    sb=0f4Tag_szj9t082e8Bc6butW;
    c_user=100021418158790;
    xs=8%3AwLNtoHNhRy9koA%3A2%3A1787020295%3A-1%3A-1%3A%3AAcxRulgHOS7-j-EXIPnB92XOMvEYaSjDMeaYvOW3eJE;
    fr=1TR1Zvw8sv0Ruknai.AWemjNpwC0f03V_ZckSoPWpFJbF_HylsudO1Ce8wElkLW-xOqXs.BqhVuw..AAA.0.0.BqhWY1.AWdojhn6aR8WDdDGMCUw_WHKuvA;
"""

def parse_cookies(cookie_string, domain=".facebook.com"):
    cookies = []
    parts = cookie_string.strip().split(";")
    for part in parts:
        if "=" in part:
            name, value = part.strip().split("=", 1)
            cookies.append({"name": name, "value": value, "domain": domain, "path": "/"})
    return cookies

def get_url_mapping(user_id):
    return {
        "intro": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_intro",
        "category": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_category",
        "personal_details": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_personal_details",
        "work": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_work",
        "education": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_education",
        "hobbies": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_activites",
        "interests": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_interests",
        "travel": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_travel",
        "links": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_links",
        "contact_info": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_contact_info",
        "names": f"https://www.facebook.com/profile.php?id={user_id}&sk=directory_names",
        "aboutyou": f"https://www.facebook.com/profile.php?id={user_id}&sk=about_details"
    }

async def scrape_facebook_profile(user_id):
    url_mapping = get_url_mapping(user_id)
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=True)
        context = await browser.new_context()

        formatted_cookies = parse_cookies(COOKIE_STRING)
        await context.add_cookies(formatted_cookies)

        page = await context.new_page()
        scraped_data = {}

        main_about_url = f"https://www.facebook.com/profile.php?id={user_id}&sk=about"
        print(f"[Үндсэн About] хуудас руу орж байна: {main_about_url}")
        
        try:
            await page.goto(main_about_url, timeout=60000)
            await page.wait_for_timeout(4000)

            full_page_text = await page.locator("body").inner_text()
            page_html = await page.content()

            username_match = re.search(r'"userVanity":"([^"]+)"', page_html)
            
            friends_match = re.search(r'([\d,\.KMB]+)\s+Friends', full_page_text, re.IGNORECASE)
            followers_match = re.search(r'([\d,\.KMB]+)\s+Followers', full_page_text, re.IGNORECASE)
            following_match = re.search(r'([\d,\.KMB]+)\s+Following', full_page_text, re.IGNORECASE)
            likes_match = re.search(r'([\d,\.KMB]+)\s+Likes', full_page_text, re.IGNORECASE)

            profile_pic = ""
            try:
                # 1. Эхлээд meta tag-ээс шалгах
                # og_image_elem = page.locator('meta[property="og:image"]')
                # if await og_image_elem.count() > 0:
                #     profile_pic = await og_image_elem.get_attribute('content')
                
                # 2. Хэрэв meta tag байхгүй бол SVG-ээс хайх 
                # (aria-label нь "Your profile" биш, тухайн хүний нэр эсвэл бусад SVG-г сонгоно)
                if not profile_pic:
                    img_elem = page.locator('svg[role="img"]:not([aria-label="Your profile"]) image').first
                    if await img_elem.count() > 0:
                        profile_pic = await img_elem.get_attribute('xlink:href')
                        if not profile_pic:
                            profile_pic = await img_elem.get_attribute('href')
            except Exception as e:
                print(f"Зураг авахад алдаа: {e}")

            stats = {}
            if friends_match:
                stats["friends"] = friends_match.group(1).replace(",", "")
            if followers_match:
                stats["followers"] = followers_match.group(1).replace(",", "")
            if following_match:
                stats["following"] = following_match.group(1).replace(",", "")
            if likes_match:
                stats["likes"] = likes_match.group(1).replace(",", "")

            if username_match:
                stats["username"] = username_match.group(1)
            else:
                stats["username"] = ""
            stats["profile_pic"] = profile_pic

            # Профайл түгжээтэй эсэхийг шалгах энгийн логик
            is_closed = "locked her profile" in full_page_text.lower() or "locked his profile" in full_page_text.lower()
            stats["is_closed"] = "Yes" if is_closed else "No"

            if stats:
                scraped_data["stats"] = stats

            active_keys = set()
            sidebar_container = page.locator('div.x1yztbdb').first
            
            if await sidebar_container.count() > 0:
                menu_items = sidebar_container.locator('a, span, div[role="link"]')
                count = await menu_items.count()
                found_texts = []
                
                for i in range(count):
                    txt = await menu_items.nth(i).inner_text()
                    if txt:
                        found_texts.append(txt.strip().lower())
                
                mapping_keywords = {
                    "intro": "intro",
                    "category": "category",
                    "personal details": "personal_details",
                    "work": "work",
                    "education": "education",
                    "hobbies": "activites",
                    "interests": "interests",
                    "travel": "travel",
                    "links": "links",
                    "contact info": "contact_info",
                    "names": "names",
                    "details about you": "aboutyou"
                }

                for key_phrase, url_key in mapping_keywords.items():
                    if any(key_phrase in ft for ft in found_texts):
                        active_keys.add(url_key)

            if not active_keys:
                active_keys = {"intro", "category", "contact_info"}

            for key in active_keys:
                if key not in url_mapping:
                    continue
                url = url_mapping[key]
                try:
                    await page.goto(url, timeout=60000)
                    await page.wait_for_timeout(3000)

                    data = await extract_data_by_section(page, key)
                    if data:
                        scraped_data[key] = data
                except Exception as e:
                    print(f"Алдаа гарлаа {key}: {e}")

        except Exception as e:
            print(f"Үндсэн хуудас руу ороход алдаа гарлаа: {e}")

        await browser.close()
        return scraped_data

async def extract_data_by_section(page, section_type):
    extracted = {}

    if section_type == "intro":
        try:
            bio_heading = page.locator('h2').filter(has_text="Bio").first
            if await bio_heading.count() > 0:
                section = bio_heading.locator("xpath=ancestor::section[1]")
                bio_locator = section.locator('div[dir="auto"][style*="text-align"]').first
                if await bio_locator.count() > 0:
                    bio = (await bio_locator.inner_text()).strip()
                    extracted["bio"] = bio if bio else None
        except Exception as e:
            print(f"Bio авахад алдаа: {e}")
    
    try:
        main = page.locator("div[role='main']")
        headings = main.locator("h2")
        result = []

        for i in range(await headings.count()):
            heading = headings.nth(i)
            title = (await heading.inner_text()).strip()
            if not title:
                continue

            section = heading.locator("xpath=ancestor::section[1]")
            if await section.count() == 0:
                continue

            texts = await section.locator('[dir="auto"]').all_inner_texts()
            clean_texts = []

            for text in texts:
                text = text.strip()
                if not text:
                    continue

                unwanted = [
                    "Shared with Public", "Shared with Friends", 
                    "Edit", "Edit details", "See more", "See less"
                ]

                if text in unwanted or text == title:
                    continue

                if text not in clean_texts:
                    clean_texts.append(text)

            if clean_texts:
                result.append({
                    "title": title,
                    "values": clean_texts
                })

        if result:
            extracted["data"] = result

    except Exception as e:
        print(f"{section_type} мэдээлэл авахад алдаа: {e}")

    return extracted

def _first_value(values):
    if not values:
        return ""
    return " | ".join(str(x).strip() for x in values if str(x).strip())

def flatten_scraped_data(scraped_data, user_id):
    # Таны хүссэн яг тэр дарааллаар багануудыг үүсгэж байна
    row = {
        "ID": str(user_id),
        "username": "",
        "profile_pic": "",
        "friends": "",
        "followers": "",
        "following": "",
        "likes": "",
        "is_closed": "No",
        "bio": "",
        "category": "",
        "location": "",
        "hometown": "",
        "birthday": "",
        "status": "",
        "family_members": "",
        "gender": "",
        "work": "",
        "education": "",
        "links": "",
        "social_media": "",
        "phone": "",
        "email": ""
    }

    stats = scraped_data.get("stats", {})
    if stats.get("username"):
        row["username"] = str(stats["username"])
    if stats.get("friends"):
        row["friends"] = str(stats["friends"])
    if stats.get("followers"):
        row["followers"] = str(stats["followers"])
    if stats.get("following"):
        row["following"] = str(stats["following"])
    if stats.get("likes"):
        row["likes"] = str(stats["likes"])
    if stats.get("is_closed"):
        row["is_closed"] = str(stats["is_closed"])
    if stats.get("profile_pic"):
        row["profile_pic"] = str(stats["profile_pic"])
    

    for section in (
        "intro", "category", "personal_details", "work", "education", "links", "contact_info"):
        section_data = scraped_data.get(section, {})
        items = section_data.get("data", [])
        if not isinstance(items, list):
            continue

        extras = []
        for item in items:
            if not isinstance(item, dict):
                continue

            title = str(item.get("title", "")).strip()
            values = item.get("values", [])
            if not isinstance(values, list):
                values = [values]

            value = _first_value(values)
            if not value:
                continue

            low = title.lower()

            if section == "intro":
                if "bio" in low:
                    row["bio"] = value
                else:
                    extras.append(f"{title}: {value}")
            elif section == "work":
                clean_work = value.split(" | ")[0]
                row["work"] = (row["work"] + " | " if row["work"] else "") + clean_work
            elif section == "education":
                row["education"] = (row["education"] + " | " if row["education"] else "") + f"{title}: {value}"
            elif section == "contact_info":
                if "email" in low:
                    row["email"] = value
                elif "phone" in low or "mobile" in low:
                    row["phone"] = value
                elif "social" in low or "instagram" in low or "twitter" in low:
                    # Social media утгуудыг хоригоор (username, platform) нь салгаж холбох
                    social_pairs = []
                    j = 0
                    while j < len(values):
                        username = values[j]
                        if j + 1 < len(values):
                            platform = values[j+1]
                            social_pairs.append(f"{username} ({platform})")
                            j += 2
                        else:
                            social_pairs.append(username)
                            j += 1
                    
                    social_str = " | ".join(social_pairs)
                    row["social_media"] = (row["social_media"] + " | " if row["social_media"] else "") + social_str
                else:
                    extras.append(f"{title}: {value}")
            elif section == "personal_details":
                if "birth" in low:
                    row["birthday"] = value
                elif "live" in low or "current" in low or "location" in low:
                    row["location"] = value.split(" | ")[0]
                elif "from" in low or "hometown" in low:
                    row["hometown"] = value.split(" | ")[0]
                elif "gender" in low or "sex" in low:
                    row["gender"] = value
                elif "status" in low:
                    row["status"] = value
                elif "family" in low:
                    family_pairs = []
                    j = 0
                    while j < len(values):
                        name = values[j]
                        if j + 1 < len(values):
                            relation = values[j+1]
                            family_pairs.append(f"{name} ({relation})")
                            j += 2
                        else:
                            family_pairs.append(name)
                            j += 1
                    
                    family_str = " | ".join(family_pairs)
                    row["family_members"] = (row["family_members"] + " | " if row["family_members"] else "") + family_str
                else:
                    extras.append(f"{title}: {value}")
            elif section == "category":
                row["category"] = value
            elif section == "links":
                clean_link = value.split(" | ")[0]
                row["links"] = (row["links"] + " | " if row["links"] else "") + clean_link
            else:
                extras.append(f"{title}: {value}")

        if section == "contact_info" and extras:
            extra_str = " | ".join(extras)
            row["social_media"] = row["social_media"] + (" | " if row["social_media"] else "") + extra_str

    return row

@app.post("/scrape")
def scrape():
    data = request.get_json(silent=True) or {}
    user_id = str(data.get("user_id", "")).strip()

    if not user_id.isdigit():
        return jsonify({"ok": False, "error": f"ID буруу байна: {user_id}"}), 400

    try:
        scraped = asyncio.run(scrape_facebook_profile(user_id))
        row = flatten_scraped_data(scraped or {}, user_id)
        return jsonify({"ok": True, "row": row})
    except Exception as e:
        return jsonify({"ok": False, "error": str(e)}), 500

@app.post("/download-excel")
def download_excel():
    data = request.get_json(silent=True) or {}
    rows = data.get("rows", [])

    if not rows:
        return jsonify({"ok": False, "error": "Мэдээлэл алга."}), 400

    headers = list(rows[0].keys())
    wb = Workbook()
    ws = wb.active
    ws.title = "Facebook Data"
    ws.append(headers)

    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)

    for i, h in enumerate(headers, 1):
        max_len = len(str(h))
        for cell in ws[get_column_letter(i)]:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(i)].width = min(max(max_len + 2, 12), 45)

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)

    return send_file(
        out,
        as_attachment=True,
        download_name="facebook_profile_data.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

HTML = r'''<!DOCTYPE html>
<html lang="mn">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>FB Multi-ID Scraper</title>
<style>
*{box-sizing:border-box;margin:0;padding:0}
body{font-family:-apple-system,BlinkMacSystemFont,"Segoe UI",sans-serif;background:#f0f2f5;min-height:100vh;padding:30px 16px}
.card{background:#fff;border-radius:16px;padding:28px;max-width:1500px;margin:auto;box-shadow:0 4px 24px rgba(0,0,0,.1)}
.logo{color:#1877f2;font-size:26px;font-weight:700;text-align:center}
.sub{color:#888;text-align:center;margin:6px 0 20px;font-size:13px}
.input-wrap{margin-bottom:15px}
textarea{width:100%;height:100px;padding:12px 16px;border:2px solid #e0e0e0;border-radius:10px;font-size:14px;outline:none;resize:vertical}
textarea:focus{border-color:#1877f2}
.main-btn{width:100%;padding:13px 24px;background:#1877f2;color:#fff;border:0;border-radius:10px;font-size:15px;font-weight:600;cursor:pointer}
.main-btn:hover{background:#1565c0}
.main-btn:disabled{background:#a9c5ee;cursor:not-allowed}
.status{margin-top:12px;color:#777;font-size:13px;min-height:20px;text-align:center}
.toolbar{display:none;gap:8px;margin-top:16px}
.tool-btn{padding:9px 14px;border:1px solid #1877f2;border-radius:8px;background:#f0f6ff;color:#1877f2;cursor:pointer;font-weight:600}
.table-wrap{display:none;margin-top:14px;border:1px solid #e2e2e2;border-radius:10px;overflow:auto;max-height:600px}
table{border-collapse:collapse;width:max-content;min-width:100%}
th{position:sticky;top:0;background:#1877f2;color:#fff;padding:10px;text-align:left;font-size:13px;white-space:nowrap}
td{padding:9px 10px;border-bottom:1px solid #eee;font-size:13px;white-space:nowrap;max-width:400px;overflow:hidden;text-overflow:ellipsis}
tr:hover td{background:#f7faff}
.copy{margin-left:5px;font-size:10px;border:1px solid rgba(255,255,255,.7);border-radius:4px;padding:2px 4px;cursor:pointer}
.empty{color:#bbb}
</style>
</head>
<body>
<div class="card">
<div class="logo">FB Multi-ID Profile Scraper</div>
<div class="sub">Олон ID нэг дор оруулан мэдээлэл цуглуулж Excel рүү татах эсвэл хуулах</div>

<div class="input-wrap">
<textarea id="userIdsInput" placeholder="ID-уудаа мөр мөрөөр эсвэл таслалаар тусгаарлаад оруулна уу...&#10;Жишээ нь:&#10;100054640671737&#10;100021418158790"></textarea>
</div>
<button class="main-btn" id="scrapeBtn" onclick="startScraping()">🔍 Бүгдийг Цуглуулах</button>

<div class="status" id="status"></div>

<div class="toolbar" id="toolbar">
<button class="tool-btn" onclick="copyAll()">📋 Бүгдийг хуулах</button>
<button class="tool-btn" onclick="downloadExcel()">📥 Excel татах</button>
</div>

<div class="table-wrap" id="tableWrap">
<table>
<thead id="thead"></thead>
<tbody id="tbody"></tbody>
</table>
</div>
</div>

<script>
let rows=[],headers=[];

function setStatus(t){
    document.getElementById("status").textContent=t;
}

async function startScraping(){
    const rawText = document.getElementById("userIdsInput").value.trim();
    if(!rawText){
        setStatus("⚠️ Хамгийн багадаа нэг ID оруулна уу.");
        return;
    }

    const ids = rawText.split(/[\n,\s]+/).map(id => id.trim()).filter(id => /^\d+$/.test(id));

    if(ids.length === 0){
        setStatus("⚠️ Хүчинтэй тоон Facebook ID олдсонгүй.");
        return;
    }

    const btn = document.getElementById("scrapeBtn");
    btn.disabled = true;
    rows = [];
    document.getElementById("toolbar").style.display = "none";
    document.getElementById("tableWrap").style.display = "none";

    for(let i = 0; i < ids.length; i++){
        const userId = ids[i];
        btn.textContent = `⏳ Цуглуулж байна... (${i+1}/${ids.length}: ID ${userId})`;
        setStatus(`[${i+1}/${ids.length}] ID: ${userId} мэдээллийг татаж байна...`);

        try {
            const res = await fetch("/scrape", {
                method: "POST",
                headers: {"Content-Type": "application/json"},
                body: JSON.stringify({user_id: userId})
            });
            const data = await res.json();
            if(data.ok){
                rows.push(data.row);
                renderTable();
                document.getElementById("toolbar").style.display = "flex";
                document.getElementById("tableWrap").style.display = "block";
            } else {
                console.error(`ID ${userId} алдаа:`, data.error);
            }
        } catch(e) {
            console.error(`ID ${userId} сүлжээний алдаа:`, e);
        }
    }

    btn.disabled = false;
    btn.textContent = "🔍 Бүгдийг Цуглуулах";
    setStatus(`✓ Нийт ${rows.length} профайлын мэдээлэл амжилттай цугларлаа.`);
}

function renderTable(){
    if(rows.length === 0) return;
    const thead = document.getElementById("thead");
    const tbody = document.getElementById("tbody");

    thead.innerHTML = "";
    tbody.innerHTML = "";
    headers = Object.keys(rows[0]);

    const trh = document.createElement("tr");
    headers.forEach(h => {
        const th = document.createElement("th");
        th.textContent = h;

        const c = document.createElement("span");
        c.className = "copy";
        c.textContent = "copy";
        c.onclick = e => {
            e.stopPropagation();
            copyColumn(h, c);
        };

        th.appendChild(c);
        trh.appendChild(th);
    });
    thead.appendChild(trh);

    rows.forEach(row => {
        const tr = document.createElement("tr");
        headers.forEach(h => {
            const td = document.createElement("td");
            const v = row[h] || "";
            td.textContent = v || "—";
            if(!v) td.className = "empty";
            tr.appendChild(td);
        });
        tbody.appendChild(tr);
    });
}

function copyText(text){
    if(navigator.clipboard && window.isSecureContext){
        navigator.clipboard.writeText(text);
    } else {
        const ta = document.createElement("textarea");
        ta.value = text;
        document.body.appendChild(ta);
        ta.select();
        document.execCommand("copy");
        ta.remove();
    }
}

function copyColumn(key, el){
    copyText(rows.map(r => r[key] || "").join("\n"));
    const old = el.textContent;
    el.textContent = "✓";
    setTimeout(() => el.textContent = old, 1200);
}

function copyAll(){
    const text = [
        headers.join("\t"),
        ...rows.map(r => headers.map(h => r[h] || "").join("\t"))
    ].join("\n");

    copyText(text);
    setStatus("✓ Бүх хүснэгтийн мэдээлэл clipboard-д хуулагдлаа.");
    setTimeout(() => setStatus(""), 2000);
}

function downloadExcel(){
    if(!rows.length) return;

    fetch("/download-excel", {
        method: "POST",
        headers: {"Content-Type": "application/json"},
        body: JSON.stringify({rows: rows})
    })
    .then(r => {
        if(!r.ok) throw new Error("Excel үүсгэхэд алдаа гарлаа.");
        return r.blob();
    })
    .then(blob => {
        const url = URL.createObjectURL(blob);
        const a = document.createElement("a");
        a.href = url;
        a.download = "facebook_profile_data.xlsx";
        document.body.appendChild(a);
        a.click();
        a.remove();
        URL.revokeObjectURL(url);
    })
    .catch(e => setStatus("❌ " + e.message));
}
</script>
</body>
</html>'''

@app.get("/")
def index():
    return render_template_string(HTML)

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5000, debug=False)